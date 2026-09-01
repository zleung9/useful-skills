#!/usr/bin/env python3
"""Batch test all API keys from an xlsx file against 5 Huawei model endpoints.
Writes results back as a new column in the xlsx.

Usage:
    python3 batch_test_keys.py <input.xlsx> [output.xlsx]

If output path omitted, appends '-tested' before .xlsx extension.
"""
import sys
import time
import json
import urllib.request
import urllib.error
import openpyxl

MODELS = [
    ("DeepSeek-V4", "http://10.26.15.52:30081/132bc0947fc64cc79e22618a60394789/v1/chat/completions", "DeepSeek-V4-Flash-w8a8-mtp"),
    ("Qwen3.6-128K", "http://10.26.15.52:30081/f6f71ef40c934f75920fb8decd6db721/v1/chat/completions", "Qwen3.6-35B-A3B"),
    ("Qwen3.6-std", "http://10.26.15.52:30081/ccb44b4cfe18439f8affd07babd0810e/v1/chat/completions", "Qwen3.6-35B-A3B"),
    ("Intern-S2", "http://10.26.15.52:30081/784fa5165cd6424fa19764c8c6b95274/v1/chat/completions", "intern-s2-preview"),
    ("GLM-5", "http://10.26.15.52:30081/7f82733149be43a1b8f26196b2202fa6/v1/chat/completions", "glm-5"),
]


def test_key(key, models=None):
    """Test one API key against all models. Returns dict {model_name: 'OK'|'HTTP401'|'ERR'} and ok_count."""
    if models is None:
        models = MODELS
    results = {}
    ok_count = 0
    for model_name, url, model_id in models:
        body = json.dumps({
            "model": model_id,
            "messages": [{"role": "user", "content": "Hi"}],
            "max_tokens": 8
        }).encode()
        req = urllib.request.Request(url, data=body, headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {key}"
        })
        try:
            with urllib.request.urlopen(req, timeout=15) as resp:
                json.loads(resp.read().decode())  # verify parseable
                results[model_name] = "OK"
                ok_count += 1
        except urllib.error.HTTPError as e:
            results[model_name] = f"HTTP{e.code}"
        except Exception:
            results[model_name] = "ERR"
        time.sleep(0.3)
    return results, ok_count


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <input.xlsx> [output.xlsx]")
        sys.exit(1)

    input_path = sys.argv[1]
    if len(sys.argv) >= 3:
        output_path = sys.argv[2]
    else:
        base, ext = input_path.rsplit(".", 1)
        output_path = f"{base}-tested.{ext}"

    wb = openpyxl.load_workbook(input_path)
    ws = wb.active

    # Find next empty column
    next_col = ws.max_column + 1

    # Add headers
    from datetime import date
    ws.cell(row=1, column=next_col, value=f"{date.today().isoformat()}测试")
    ws.cell(row=2, column=next_col, value="5模型连通测试")
    ws.cell(row=3, column=next_col, value="有效性(≥4连通)")

    # Find API key column: look for header row with 'API KEY' (not '名称'),
    # then key values are in the next column
    key_col = None
    for col_idx in range(1, ws.max_column + 1):
        for row_idx in range(1, min(5, ws.max_row + 1)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val and "API KEY" in str(val).upper() and "名称" not in str(val):
                key_col = col_idx + 1  # key value is next column
                break
        if key_col:
            break

    if not key_col:
        # Fallback: find first column with 'sk-' values
        for row_idx in range(1, min(10, ws.max_row + 1)):
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val and str(val).startswith("sk-"):
                    key_col = col_idx
                    break
            if key_col:
                break

    if not key_col:
        print("ERROR: Could not find API key column in xlsx")
        sys.exit(1)

    print(f"Using key column: {key_col}")
    total = 0
    valid = 0

    for row_idx in range(3, ws.max_row + 1):
        key_val = ws.cell(row=row_idx, column=key_col).value
        if not key_val or not str(key_val).startswith("sk-"):
            continue

        total += 1
        key_str = str(key_val)
        name = ws.cell(row=row_idx, column=2).value or f"Row{row_idx}"
        key_short = key_str[:8] + "..." + key_str[-4:]

        model_results, ok_count = test_key(key_str)
        status = "有效" if ok_count >= 4 else "无效"
        if ok_count >= 4:
            valid += 1

        model_str = ", ".join(f"{m}:{v}" for m, v in model_results.items())
        cell_value = f"{status}({ok_count}/5) {model_str}"
        ws.cell(row=row_idx, column=next_col, value=cell_value)

        print(f"[{total}] {name} ({key_short}): {ok_count}/5 - {status}")

    wb.save(output_path)
    print(f"\nSummary: {valid} 有效 / {total - valid} 无效 / {total} total")
    print(f"Saved to: {output_path}")


if __name__ == "__main__":
    main()
