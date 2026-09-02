#!/usr/bin/env python3
"""Shared helpers for academic-integrity screening scripts."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sys
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TABLE_EXTS = {".csv", ".tsv", ".xlsx", ".xlsm"}
TEXT_EXTS = {".txt", ".md", ".ris", ".bib", ".nbib", ".enw"}
DOC_EXTS = {".pdf", ".docx", ".doc", ".tex", ".rtf"}

SEVERITY_ORDER = {
    "INFO": 0,
    "LOW": 1,
    "MEDIUM": 2,
    "HIGH": 3,
    "CRITICAL": 4,
}


def add_common_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--format", choices=["markdown", "json"], default="markdown")
    parser.add_argument("--output", help="Optional output file")


def make_finding(
    *,
    method: str,
    severity: str,
    description: str,
    location: str = "",
    evidence_files: Optional[Sequence[str]] = None,
    statistics: Optional[Dict[str, object]] = None,
    recommendation: str = "",
) -> Dict[str, object]:
    return {
        "method": method,
        "severity": severity.upper(),
        "location": location,
        "description": description,
        "statistics": statistics or {},
        "evidence_files": list(evidence_files or []),
        "recommendation": recommendation,
    }


def risk_level(findings: Sequence[Dict[str, object]]) -> str:
    counts = {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "CRITICAL": 0}
    for finding in findings:
        severity = str(finding.get("severity", "INFO")).upper()
        if severity in counts:
            counts[severity] += 1
    if counts["CRITICAL"] >= 2 or counts["HIGH"] >= 5:
        return "BLACK investigation threshold"
    if counts["CRITICAL"] >= 1 or counts["HIGH"] >= 2:
        return "RED severe"
    if counts["HIGH"] >= 1 or counts["MEDIUM"] >= 3:
        return "ORANGE high"
    if counts["MEDIUM"] >= 1 or counts["LOW"] >= 1:
        return "YELLOW moderate"
    return "GREEN low"


def make_result(
    *,
    tool: str,
    input_value: object,
    findings: Sequence[Dict[str, object]],
    evidence_files: Optional[Sequence[str]] = None,
    limitations: Optional[Sequence[str]] = None,
    metadata: Optional[Dict[str, object]] = None,
) -> Dict[str, object]:
    return {
        "tool": tool,
        "input": input_value,
        "risk_level": risk_level(findings),
        "findings": list(findings),
        "evidence_files": list(evidence_files or []),
        "limitations": list(limitations or []),
        "metadata": metadata or {},
    }


def format_value(value: object) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def result_to_markdown(result: Dict[str, object]) -> str:
    lines = [
        f"# {result.get('tool', 'Integrity Screen')}",
        "",
        f"- Input: `{format_value(result.get('input', ''))}`",
        f"- Risk level: {result.get('risk_level', 'UNKNOWN')}",
        "",
        "## Findings",
        "",
    ]
    findings = result.get("findings", [])
    if not findings:
        lines.append("No screening flags were detected.")
    else:
        for idx, finding in enumerate(findings, 1):
            lines.extend(
                [
                    f"### Finding {idx}: {finding.get('method', '')}",
                    "",
                    f"- Severity: {finding.get('severity', '')}",
                    f"- Location: {finding.get('location', '') or 'not specified'}",
                    f"- Description: {finding.get('description', '')}",
                ]
            )
            evidence = finding.get("evidence_files") or []
            if evidence:
                lines.append(f"- Evidence files: {', '.join(f'`{item}`' for item in evidence)}")
            recommendation = finding.get("recommendation")
            if recommendation:
                lines.append(f"- Recommendation: {recommendation}")
            statistics = finding.get("statistics") or {}
            if statistics:
                lines.append("- Statistics:")
                for key, value in statistics.items():
                    lines.append(f"  - {key}: {format_value(value)}")
            lines.append("")
    limitations = result.get("limitations") or []
    if limitations:
        lines.extend(["## Limitations", ""])
        for item in limitations:
            lines.append(f"- {item}")
        lines.append("")
    return "\n".join(lines)


def emit_result(result: Dict[str, object], fmt: str, output: Optional[str]) -> str:
    text = json.dumps(result, ensure_ascii=False, indent=2) if fmt == "json" else result_to_markdown(result)
    if output:
        Path(output).write_text(text, encoding="utf-8")
    else:
        print(text)
    return text


def iter_files(paths: Sequence[str], exts: Optional[Set[str]] = None) -> List[Path]:
    files: List[Path] = []
    for raw in paths:
        path = Path(raw)
        if path.is_file():
            if exts is None or path.suffix.lower() in exts:
                files.append(path)
        elif path.is_dir():
            for child in path.rglob("*"):
                if child.is_file() and (exts is None or child.suffix.lower() in exts):
                    files.append(child)
    return sorted(files)


def sha256_file(path: Path, block_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(block_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def sniff_csv(path: str, delimiter: Optional[str] = None) -> List[Dict[str, object]]:
    if path == "-":
        content = sys.stdin.read()
        sample = content[:4096]
        if delimiter is not None:
            return list(csv.DictReader(content.splitlines(), delimiter=delimiter))
        try:
            dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(content.splitlines(), dialect=dialect))

    with open(path, newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        if delimiter is not None:
            return list(csv.DictReader(handle, delimiter=delimiter))
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel_tab if path.endswith(".tsv") else csv.excel
        return list(csv.DictReader(handle, dialect=dialect))


def load_table(path: str, delimiter: Optional[str] = None) -> List[Dict[str, object]]:
    suffix = Path(path).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("XLSX input requires openpyxl: python3 -m pip install openpyxl") from exc
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return []
        columns = [str(cell).strip() if cell is not None else "" for cell in header]
        rows: List[Dict[str, object]] = []
        for raw_row in rows_iter:
            row = {}
            for idx, cell in enumerate(raw_row):
                if idx < len(columns) and columns[idx]:
                    row[columns[idx]] = cell
            rows.append(row)
        return rows
    return sniff_csv(path, delimiter)


def parse_float(value: object) -> Optional[float]:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def numeric_columns(rows: Sequence[Dict[str, object]], min_n: int = 2) -> Dict[str, List[float]]:
    values: Dict[str, List[float]] = {}
    columns = sorted({str(key) for row in rows for key in row.keys()})
    for column in columns:
        nums = [parse_float(row.get(column)) for row in rows]
        clean = [num for num in nums if num is not None]
        if len(clean) >= min_n:
            values[column] = clean
    return values


def mean(values: Sequence[float]) -> float:
    return sum(values) / len(values)


def sample_sd(values: Sequence[float]) -> float:
    if len(values) < 2:
        return 0.0
    mu = mean(values)
    return math.sqrt(sum((value - mu) ** 2 for value in values) / (len(values) - 1))


def hamming(a: str, b: str) -> int:
    return sum(ch1 != ch2 for ch1, ch2 in zip(a, b))
