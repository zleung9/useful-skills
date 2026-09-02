#!/usr/bin/env python3
"""Numeric academic-integrity screen inspired by Geng-style checks.

This script is a screening helper, not a misconduct detector. It reads numeric
columns from CSV/TSV/stdin and, when openpyxl is installed, XLSX files.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple


CHI2_CRITICAL_DF9 = {
    "p<0.05": 16.919,
    "p<0.01": 21.666,
    "p<0.001": 27.877,
}


@dataclass
class Finding:
    table: str
    column: str
    method: str
    severity: str
    description: str
    statistics: Dict[str, object]


def clean_number_text(value: object) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    if not re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", text):
        return None
    return text


def to_decimal(text: str) -> Optional[Decimal]:
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def decimal_key(value: Decimal) -> str:
    return format(value.normalize(), "f")


def last_reported_digit(text: str) -> Optional[int]:
    mantissa = re.split(r"[eE]", text)[0]
    digits = re.sub(r"\D", "", mantissa)
    if not digits:
        return None
    return int(digits[-1])


def decimal_prefix(text: str, places: int) -> Optional[str]:
    mantissa = re.split(r"[eE]", text)[0]
    if "." not in mantissa:
        return None
    frac = mantissa.split(".", 1)[1]
    frac = re.sub(r"\D", "", frac)
    if len(frac) < places:
        return None
    return frac[:places]


def load_csv_rows(path: str, delimiter: Optional[str]) -> Dict[str, List[Dict[str, object]]]:
    if path == "-":
        content = sys.stdin.read()
        sample = content[:4096]
        if delimiter is not None:
            rows = list(csv.DictReader(content.splitlines(), delimiter=delimiter))
        else:
            try:
                dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.DictReader(content.splitlines(), dialect=dialect))
        return {"stdin": rows}

    with open(path, newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        if delimiter is not None:
            rows = list(csv.DictReader(handle, delimiter=delimiter))
        else:
            try:
                dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                dialect = csv.excel_tab if path.endswith(".tsv") else csv.excel
            rows = list(csv.DictReader(handle, dialect=dialect))
    return {Path(path).name: rows}


def load_xlsx_rows(path: str) -> Dict[str, List[Dict[str, object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("XLSX input requires openpyxl: python3 -m pip install openpyxl") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    tables: Dict[str, List[Dict[str, object]]] = {}
    for sheet in workbook.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        columns = [str(cell).strip() if cell is not None else "" for cell in header]
        rows = []
        for raw_row in rows_iter:
            row = {}
            for idx, cell in enumerate(raw_row):
                if idx < len(columns) and columns[idx]:
                    row[columns[idx]] = cell
            rows.append(row)
        tables[sheet.title] = rows
    return tables


def load_tables(path: str, delimiter: Optional[str]) -> Dict[str, List[Dict[str, object]]]:
    suffix = Path(path).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return load_xlsx_rows(path)
    return load_csv_rows(path, delimiter)


def numeric_columns(rows: Sequence[Dict[str, object]], min_n: int) -> Dict[str, List[Tuple[str, Decimal]]]:
    values: Dict[str, List[Tuple[str, Decimal]]] = defaultdict(list)
    for row in rows:
        for column, raw in row.items():
            text = clean_number_text(raw)
            if text is None:
                continue
            number = to_decimal(text)
            if number is None:
                continue
            values[column].append((text, number))
    return {column: vals for column, vals in values.items() if len(vals) >= min_n}


def chi_square_label(statistic: float) -> str:
    label = "not significant at p<0.05"
    for name, critical in CHI2_CRITICAL_DF9.items():
        if statistic >= critical:
            label = name
    return label


def check_tail_distribution(table: str, column: str, vals: Sequence[Tuple[str, Decimal]]) -> Optional[Finding]:
    digits = [last_reported_digit(text) for text, _ in vals]
    digits = [digit for digit in digits if digit is not None]
    if len(digits) < 10:
        return None
    observed = Counter(digits)
    expected = len(digits) / 10.0
    chi2 = sum(((observed.get(d, 0) - expected) ** 2) / expected for d in range(10))
    cramers_v = math.sqrt(chi2 / (len(digits) * 9))
    max_digit, max_count = observed.most_common(1)[0]
    deviation_ratio = max_count / expected if expected else 0
    significant = chi2 >= CHI2_CRITICAL_DF9["p<0.05"]
    size_ok_for_effect_rules = len(digits) >= 20
    flagged = significant or (size_ok_for_effect_rules and (cramers_v >= 0.3 or deviation_ratio >= 2.0))
    if not flagged:
        return None
    severity = "HIGH" if chi2 >= CHI2_CRITICAL_DF9["p<0.01"] or cramers_v >= 0.5 or deviation_ratio >= 3.0 else "MEDIUM"
    return Finding(
        table=table,
        column=column,
        method="terminal_digit_distribution",
        severity=severity,
        description="Last reported digits deviate from an approximately uniform distribution.",
        statistics={
            "n": len(digits),
            "chi_square_df9": round(chi2, 3),
            "threshold_label": chi_square_label(chi2),
            "cramers_v": round(cramers_v, 3),
            "most_common_digit": max_digit,
            "most_common_count": max_count,
            "expected_per_digit": round(expected, 2),
            "deviation_ratio": round(deviation_ratio, 2),
            "counts": {str(d): observed.get(d, 0) for d in range(10)},
        },
    )


def check_decimal_consistency(
    table: str, column: str, vals: Sequence[Tuple[str, Decimal]], places: int
) -> Optional[Finding]:
    prefixes = [decimal_prefix(text, places) for text, _ in vals]
    prefixes = [prefix for prefix in prefixes if prefix is not None]
    if len(prefixes) < 10:
        return None
    counts = Counter(prefixes)
    repeated_groups = sum(1 for count in counts.values() if count >= 2)
    repeated_items = sum(count for count in counts.values() if count >= 2)
    max_prefix, max_count = counts.most_common(1)[0]
    repeat_rate = repeated_items / len(prefixes)
    flagged = repeated_groups > 5 or max_count >= 3 or repeat_rate > 0.15
    if not flagged:
        return None
    severity = "HIGH" if max_count >= 4 or repeat_rate > 0.25 else ("MEDIUM" if max_count >= 3 else "LOW")
    return Finding(
        table=table,
        column=column,
        method=f"decimal_prefix_{places}",
        severity=severity,
        description=f"Decimal prefixes with {places} places repeat unusually often.",
        statistics={
            "n_with_decimal_prefix": len(prefixes),
            "repeated_groups": repeated_groups,
            "repeated_items": repeated_items,
            "repeat_rate": round(repeat_rate, 4),
            "max_prefix": max_prefix,
            "max_count": max_count,
            "top_repeats": dict(counts.most_common(10)),
        },
    )


def check_exact_duplicates(table: str, column: str, vals: Sequence[Tuple[str, Decimal]]) -> Optional[Finding]:
    keys = [decimal_key(number) for _, number in vals]
    counts = Counter(keys)
    duplicate_values = sum(1 for count in counts.values() if count >= 2)
    max_value, max_count = counts.most_common(1)[0]
    flagged = duplicate_values > 5 or max_count >= 3
    if not flagged:
        return None
    severity = "HIGH" if max_count >= 4 else ("MEDIUM" if max_count >= 3 else "LOW")
    return Finding(
        table=table,
        column=column,
        method="exact_value_duplication",
        severity=severity,
        description="Exact numeric values repeat more often than expected for independent continuous measurements.",
        statistics={
            "n": len(keys),
            "duplicate_values": duplicate_values,
            "max_value": max_value,
            "max_count": max_count,
            "top_repeats": dict(counts.most_common(10)),
        },
    )


def score_findings(findings: Sequence[Finding]) -> Tuple[int, str]:
    score = len(findings) * 3
    method_counts = Counter(finding.method for finding in findings)
    for count in method_counts.values():
        if count >= 3:
            score += 10
    if len(method_counts) >= 2:
        score += 5 * (len(method_counts) - 1)
    if any("4b" in finding.column.lower() or "core" in finding.column.lower() for finding in findings):
        score += 10
    if any(int(finding.statistics.get("max_count", 0)) >= 3 for finding in findings):
        score += 10
    score = min(score, 100)
    if score >= 81:
        level = "BLACK formal-investigation threshold"
    elif score >= 61:
        level = "RED severe"
    elif score >= 31:
        level = "ORANGE high"
    elif score > 0:
        level = "YELLOW moderate"
    else:
        level = "GREEN low"
    return score, level


def run_screen(path: str, min_n: int, decimal_places: Sequence[int], delimiter: Optional[str]) -> Dict[str, object]:
    tables = load_tables(path, delimiter)
    findings: List[Finding] = []
    screened_columns = 0

    for table_name, rows in tables.items():
        for column, vals in numeric_columns(rows, min_n=min_n).items():
            screened_columns += 1
            tail = check_tail_distribution(table_name, column, vals)
            if tail:
                findings.append(tail)
            for places in decimal_places:
                decimal_finding = check_decimal_consistency(table_name, column, vals, places)
                if decimal_finding:
                    findings.append(decimal_finding)
            duplicate = check_exact_duplicates(table_name, column, vals)
            if duplicate:
                findings.append(duplicate)

    score, level = score_findings(findings)
    evidence_files = [] if path == "-" else [path]
    return {
        "tool": "geng_numeric_screen",
        "input": path,
        "tables": list(tables.keys()),
        "screened_columns": screened_columns,
        "risk_score": score,
        "risk_level": level,
        "findings": [
            {
                **asdict(finding),
                "location": f"{finding.table} ({finding.column})",
                "evidence_files": evidence_files,
            }
            for finding in findings
        ],
        "evidence_files": evidence_files,
        "limitations": [
            "Screening result only. Interpret with domain context and source records.",
            "Numeric anomalies are not misconduct findings; check measurement resolution, discreteness, rounding, and experimental design.",
        ],
        "metadata": {
            "min_n": min_n,
            "decimal_places": list(decimal_places),
            "tables": list(tables.keys()),
            "screened_columns": screened_columns,
        },
        "disclaimer": "Screening result only. Interpret with domain context and source records.",
    }


def to_markdown(result: Dict[str, object]) -> str:
    lines = [
        "# Geng Numeric Integrity Screen",
        "",
        f"- Input: `{result['input']}`",
        f"- Tables: {', '.join(result['tables']) if result['tables'] else 'none'}",
        f"- Screened numeric columns: {result['screened_columns']}",
        f"- Risk score: {result['risk_score']}/100",
        f"- Risk level: {result['risk_level']}",
        "",
        "## Findings",
        "",
    ]
    findings = result["findings"]
    if not findings:
        lines.append("No numeric screening flags were detected.")
    else:
        for idx, finding in enumerate(findings, 1):
            lines.extend(
                [
                    f"### Finding {idx}: {finding['method']}",
                    "",
                    f"- Table: `{finding['table']}`",
                    f"- Column: `{finding['column']}`",
                    f"- Severity: {finding['severity']}",
                    f"- Description: {finding['description']}",
                    "- Statistics:",
                ]
            )
            for key, value in finding["statistics"].items():
                if isinstance(value, dict) and len(value) > 6:
                    value = json.dumps(value, ensure_ascii=False)
                lines.append(f"  - {key}: {value}")
            lines.append("")
    lines.extend(
        [
            "## Interpretation Limit",
            "",
            str(result["disclaimer"]),
            "Flags are not misconduct findings. Check source data, measurement resolution, experimental design, and legitimate discreteness before escalating.",
            "",
        ]
    )
    return "\n".join(lines)


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Screen numeric tables for academic-integrity red flags.")
    parser.add_argument("input", help="CSV/TSV/XLSX path, or '-' for CSV from stdin")
    parser.add_argument("--format", choices=["markdown", "json"], default="markdown")
    parser.add_argument("--output", help="Optional output file")
    parser.add_argument("--min-n", type=int, default=20, help="Minimum numeric values per column")
    parser.add_argument("--decimal-places", default="2,3", help="Comma-separated decimal prefix lengths")
    parser.add_argument("--delimiter", help="CSV delimiter override")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    decimal_places = [int(item.strip()) for item in args.decimal_places.split(",") if item.strip()]
    result = run_screen(args.input, min_n=args.min_n, decimal_places=decimal_places, delimiter=args.delimiter)
    output = json.dumps(result, ensure_ascii=False, indent=2) if args.format == "json" else to_markdown(result)
    if args.output:
        Path(args.output).write_text(output, encoding="utf-8")
    else:
        print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
